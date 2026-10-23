from pathlib import Path
import io

from docx import Document as DocxDocument
from pypdf import PdfReader
import pymupdf
import pytesseract
from PIL import Image


# Configure the Tesseract OCR executable for Windows.
# pytesseract.pytesseract.tesseract_cmd = (
#     r"C:\Program Files\Tesseract-OCR\tesseract.exe"
# )


def extract_text_from_file(file_path: str, file_type: str) -> str:
    """
    Extract text from PDF, DOCX, and TXT files.
    Use OCR as a fallback for scanned PDF pages.
    """

    path = Path(file_path)
    extension = file_type.lower().lstrip(".")

    if not path.exists() or not path.is_file():
        raise FileNotFoundError("Document file was not found")

    if extension == "pdf":
        reader = PdfReader(str(path))

        if reader.is_encrypted:
            raise ValueError(
                "Password-protected PDFs are not supported"
            )

        pages_text = []

        # Open the PDF once and reuse it for OCR when necessary.
        with pymupdf.open(str(path)) as pdf_document:
            for page_index, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""

                # Run OCR when the page has little extractable text.
                if len(page_text.strip()) < 30:
                    try:
                        pdf_page = pdf_document.load_page(page_index)

                        pixmap = pdf_page.get_pixmap(
                            matrix=pymupdf.Matrix(2, 2),
                            alpha=False,
                        )

                        image = Image.open(
                            io.BytesIO(pixmap.tobytes("png"))
                        )

                        page_text = pytesseract.image_to_string(
                            image,
                            lang="eng",
                        )

                    except Exception as exc:
                        raise ValueError(
                            "OCR failed. Check that Tesseract OCR "
                            "is installed and the PDF is readable."
                        ) from exc

                if page_text.strip():
                    pages_text.append(page_text.strip())

        extracted_text = "\n\n".join(pages_text)

    elif extension == "docx":
        doc = DocxDocument(str(path))
        parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)

        # Extract text from tables as well.
        for table in doc.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    parts.append(" | ".join(cells))

        extracted_text = "\n".join(parts)

    elif extension == "txt":
        extracted_text = path.read_text(
            encoding="utf-8-sig",
            errors="replace",
        )

    else:
        raise ValueError("Unsupported document type")

    extracted_text = extracted_text.strip()

    if not extracted_text:
        raise ValueError(
            "No readable text found. The document may be empty, "
            "scanned, or unreadable."
        )

    return extracted_text
